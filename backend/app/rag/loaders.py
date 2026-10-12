"""Document loaders: PDF, TXT, Markdown, DOCX -> list of pages."""
import re
from dataclasses import dataclass
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".markdown", ".docx"}


@dataclass
class Page:
    text: str
    page: int | None  # 1-based; None when the format has no pages


def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r" ?\n ?", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def load_document(path: Path) -> list[Page]:
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type '{ext}'. Supported: {sorted(SUPPORTED_EXTENSIONS)}")
    if ext == ".pdf":
        pages = _load_pdf(path)
    elif ext == ".docx":
        pages = _load_docx(path)
    else:
        pages = [Page(path.read_text(encoding="utf-8", errors="replace"), None)]
    cleaned = [Page(clean_text(p.text), p.page) for p in pages]
    return [p for p in cleaned if p.text]


def _load_pdf(path: Path) -> list[Page]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return [Page(page.extract_text() or "", i + 1) for i, page in enumerate(reader.pages)]


def _load_docx(path: Path) -> list[Page]:
    from docx import Document

    doc = Document(str(path))
    lines = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            lines.append("")
        elif para.style is not None and para.style.name.lower().startswith("heading"):
            level = re.search(r"\d+", para.style.name)
            lines.append("#" * (int(level.group()) if level else 1) + " " + text)
        else:
            lines.append(text)
    for table in doc.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return [Page("\n".join(lines), None)]
