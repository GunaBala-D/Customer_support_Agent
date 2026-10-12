import pytest
from docx import Document

from app.rag.loaders import load_document
from tests.pdfwriter import write_pdf


def test_txt_and_markdown(tmp_path):
    (tmp_path / "a.txt").write_text("Hello\n\n\n\nWorld", encoding="utf-8")
    (tmp_path / "b.md").write_text("# T\ntext", encoding="utf-8")
    assert load_document(tmp_path / "a.txt")[0].text == "Hello\n\nWorld"
    assert load_document(tmp_path / "b.md")[0].page is None


def test_pdf_keeps_page_numbers(tmp_path):
    write_pdf(tmp_path / "p.pdf", [["First page about shipping."], ["Second page about returns."]])
    pages = load_document(tmp_path / "p.pdf")
    assert [p.page for p in pages] == [1, 2]
    assert "shipping" in pages[0].text and "returns" in pages[1].text


def test_docx_headings_and_tables(tmp_path):
    d = Document()
    d.add_heading("Warranty", level=1)
    d.add_paragraph("Twelve months.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Storage", "36 months"
    d.save(str(tmp_path / "w.docx"))
    text = load_document(tmp_path / "w.docx")[0].text
    assert "# Warranty" in text and "Twelve months." in text and "Storage | 36 months" in text


def test_unsupported_extension(tmp_path):
    (tmp_path / "x.exe").write_bytes(b"MZ")
    with pytest.raises(ValueError):
        load_document(tmp_path / "x.exe")
