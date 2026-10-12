"""Regenerates the binary sample documents (PDF + DOCX) in data/documents.

Run once: python scripts/make_sample_docs.py   (the generated files are committed)
"""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))
from tests.pdfwriter import write_pdf  # noqa: E402

OUT = ROOT / "data" / "documents"

SHIPPING_PAGES = [
    [
        "Acme Electronics Shipping Policy",
        "",
        "Processing time",
        "Orders are processed within 1 business day. Orders placed after 2:00 PM Eastern Time, or on weekends",
        "and holidays, are processed the next business day.",
        "",
        "Domestic shipping options",
        "Standard shipping takes 3 to 5 business days. Express shipping takes 1 to 2 business days.",
        "Orders over $50 qualify for free standard shipping. Standard shipping on smaller orders costs $5.99.",
        "Express shipping costs $14.99. Premium members receive free express shipping.",
    ],
    [
        "Tracking and delivery",
        "",
        "Tracking",
        "A tracking number is emailed when your order ships. Tracking can take up to 24 hours to show updates.",
        "We ship with UPS, FedEx, USPS and DHL.",
        "",
        "International shipping",
        "International orders ship with DHL and take 7 to 14 business days. Customers are responsible for",
        "import duties and taxes charged by their country.",
    ],
    [
        "Address changes and lost packages",
        "",
        "Changing the delivery address",
        "The delivery address can only be changed while the order status is Pending or Processing.",
        "",
        "Lost or missing packages",
        "If tracking shows delivered but you cannot find the package, contact support within 14 days of the",
        "delivery date. We will open a carrier investigation and send a replacement or refund if the package",
        "is confirmed lost.",
    ],
]


def make_docx(path: Path) -> None:
    from docx import Document

    d = Document()
    d.add_heading("Acme Electronics Warranty Policy", level=1)
    d.add_heading("Coverage period", level=2)
    d.add_paragraph("All Acme Electronics products include a 12-month limited warranty from the delivery date. "
                    "Premium members receive 24 months on eligible products. Some products list a longer "
                    "period on their product page, for example 36 months for storage devices.")
    d.add_heading("What is covered", level=2)
    d.add_paragraph("The warranty covers manufacturing defects and hardware failures under normal use. "
                    "We repair or replace the product, or issue a refund if no replacement is available.")
    d.add_heading("What is not covered", level=2)
    d.add_paragraph("The warranty does not cover water damage, accidental drops, cracked screens, "
                    "unauthorized repairs, battery wear after 12 months, or cosmetic wear.")
    d.add_heading("Making a warranty claim", level=2)
    d.add_paragraph("Contact support with your order number, the product serial number and photos or a short "
                    "video of the problem. Claims are reviewed within 3 business days. Approved claims are "
                    "shipped a prepaid return label.")
    t = d.add_table(rows=3, cols=2)
    for row, (a, b) in zip(t.rows, [("Product category", "Warranty"), ("Storage devices", "36 months"), ("Smart home locks", "24 months")], strict=True):
        row.cells[0].text, row.cells[1].text = a, b
    d.save(str(path))


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    write_pdf(OUT / "shipping_policy.pdf", SHIPPING_PAGES)
    make_docx(OUT / "warranty_policy.docx")
    print("wrote", sorted(p.name for p in OUT.iterdir()))
